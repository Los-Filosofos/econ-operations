"""Export the tables of the ENTREGABLES package to Excel-ready CSV and XLSX.

The Markdown of each deliverable stays the editable source (RNF-02); this script copies its
table, without adding rows or values, so that every spreadsheet in `ENTREGABLES/` shows the
same matrix as the deliverable it is named after. It fixes three defects of the hand-copied
package: the CSV had no BOM (Excel showed «SecciÃ³n» instead of «Sección»), the workbook had
no column widths, header style or filters, and its content did not match its file name.

Outputs, all written to `ENTREGABLES/`:

- `1-Matriz-de-Mapeo-de-Campos.csv` and `.xlsx`: the consolidated mapping of deliverable 1,
  plus the 17 detail sheets taken from `docs/equivalencias-prisma-startrack.md` and
  `docs/matriz-requisitos-entregables.md` (the same tables as `output/matrices`).
- `1-Matriz-de-Mapeo-de-Campos-Detalle-Formulario.csv`: the complete task-form matrix.
- `2-Matriz-de-Responsabilidades-RACI.csv` and `.xlsx`: the RACI by management area of
  deliverable 2, with the proposed RACI by decision as the detail sheet.
- `2-Matriz-de-Responsabilidades-RACI-Detalle.csv`: that proposed RACI on its own.
- `MANIFEST-TABLAS.md` and `manifest-tablas.json`: SHA-256 of every source and every output.

CSV is UTF-8 with BOM and CRLF records so Excel reads the accents on a double click; the
field separator stays the comma of RFC 4180, because the CSV is for automated ingest and the
XLSX is the file to open in Excel. Reuses the parser of `exportar_matrices.py`.

Standard library plus openpyxl 3.1.5, pinned in the `dev` group of `apps/api`:

    uv run --project apps/api python scripts/docs/exportar_entregables.py
    uv run --no-project --with openpyxl==3.1.5 python scripts/docs/exportar_entregables.py --check

`--check` rebuilds everything in memory (keeping the date stored in the manifest) and fails
when any CSV, workbook or manifest differs from the checked-in files.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import sys
import zipfile
from dataclasses import dataclass
from datetime import UTC, date, datetime, time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import exportar_matrices as matrices  # noqa: E402

MATRIX_SOURCES = matrices.SOURCES
OPENPYXL_VERSION = matrices.OPENPYXL_VERSION
ZIP_EPOCH = matrices.ZIP_EPOCH
Table = matrices.Table
TableSpec = matrices.TableSpec

ROOT = Path(__file__).resolve().parents[2]
OUTPUT_DIR = Path("ENTREGABLES")
MAPPING_DOC = OUTPUT_DIR / "1-Matriz-de-Mapeo-de-Campos.md"
RACI_DOC = OUTPUT_DIR / "2-Matriz-de-Responsabilidades-RACI.md"
MANIFEST_JSON = OUTPUT_DIR / "manifest-tablas.json"
MANIFEST_MD = OUTPUT_DIR / "MANIFEST-TABLAS.md"
REPRODUCE = (
    f"uv run --no-project --with openpyxl=={OPENPYXL_VERSION} "
    "python scripts/docs/exportar_entregables.py"
)

# Brand blue of ECON (`app/dashboard/theme.py`): identity of the header band, never a state.
HEADER_FILL = "144F81"
HEADER_FONT = "FFFFFF"
MIN_WIDTH = 12
MAX_WIDTH = 58
HEADER_WIDTH = 22


@dataclass(frozen=True)
class Deliverable:
    """One numbered deliverable: its own table, its detail tables and its file names."""

    stem: str
    title: str
    document: Path
    heading: str
    sheet: str
    detail_csv: str
    detail_slugs: tuple[str, ...]
    extra_sheets: tuple[str, ...]


DELIVERABLES: tuple[Deliverable, ...] = (
    Deliverable(
        stem="1-Matriz-de-Mapeo-de-Campos",
        title="Entregable 1 · Matriz de Mapeo de Campos",
        document=MAPPING_DOC,
        heading="## 1. Mapeo Consolidado de Correspondencia",
        sheet="Mapeo consolidado",
        detail_csv="1-Matriz-de-Mapeo-de-Campos-Detalle-Formulario.csv",
        detail_slugs=("formulario-tarea-completo",),
        # Every canonical matrix travels with deliverable 1, as the previous workbook did.
        extra_sheets=("*",),
    ),
    Deliverable(
        stem="2-Matriz-de-Responsabilidades-RACI",
        title="Entregable 2 · Matriz de Responsabilidades (RACI)",
        document=RACI_DOC,
        heading="## 1. Tabla Estructurada RACI",
        sheet="RACI por gerencia",
        detail_csv="2-Matriz-de-Responsabilidades-RACI-Detalle.csv",
        detail_slugs=("raci-propuesta",),
        extra_sheets=("raci-propuesta",),
    ),
)


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def unemphasise(text: str) -> str:
    """Drop the single-asterisk emphasis that `plain()` leaves, e.g. «*(Ninguno)*»."""
    return re.sub(r"\*([^*\n]+)\*", r"\1", text)


def deliverable_table(spec: TableSpec, document: str, source: Path) -> Table:
    blocks = matrices.tables_after(document.splitlines(), spec.heading, source)
    if not blocks:
        raise SystemExit(f"«{spec.heading}» no tiene ninguna tabla en {source}.")
    table = matrices.parse_table(spec, blocks[0])
    header = [unemphasise(cell) for cell in table.header]
    rows = [[unemphasise(cell) for cell in row] for row in table.rows]
    return Table(spec, header, rows, table.section)


def csv_bytes(table: Table) -> bytes:
    """UTF-8 with BOM and CRLF records: Excel keeps the accents on a double click."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(table.header)
    writer.writerows(table.rows)
    return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")


def widest_line(values: list[str]) -> int:
    return max((len(line) for value in values for line in value.split("\n")), default=0)


def column_width(header: str, body: list[str]) -> float:
    """The body decides the width; the header only widens it a little, because it wraps."""
    longest = max(widest_line(body), min(widest_line([header]), HEADER_WIDTH))
    return float(min(MAX_WIDTH, max(MIN_WIDTH, longest + 2)))


def write_sheet(book, table: Table, title: str) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheet = book.create_sheet(title[:31])
    sheet.append(table.header)
    for row in table.rows:
        sheet.append(row)
    header_font = Font(bold=True, color=HEADER_FONT)
    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    header_alignment = Alignment(vertical="center", wrap_text=True)
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    body_alignment = Alignment(vertical="top", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = body_alignment
    for index in range(1, len(table.header) + 1):
        body = [row[index - 1] for row in table.rows]
        sheet.column_dimensions[get_column_letter(index)].width = column_width(
            table.header[index - 1], body
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False


def workbook_bytes(tables: list[tuple[Table, str]], title: str, generated_on: date) -> bytes:
    try:
        import openpyxl
        from openpyxl.writer.excel import ExcelWriter
    except ImportError as error:  # pragma: no cover - depends on the invocation
        raise SystemExit(f"openpyxl no está disponible; ejecutar con: {REPRODUCE}") from error
    if openpyxl.__version__ != OPENPYXL_VERSION:
        raise SystemExit(
            f"openpyxl {openpyxl.__version__} distinto de {OPENPYXL_VERSION}: el libro no "
            f"sería reproducible byte a byte. Ejecutar con: {REPRODUCE}"
        )
    book = openpyxl.Workbook()
    book.remove(book.active)
    for table, sheet_title in tables:
        write_sheet(book, table, sheet_title)
    # openpyxl serialises a naive timestamp as UTC midnight; the date is the only input.
    stamp = datetime.combine(generated_on, time(0, 0))
    book.properties.title = title
    book.properties.creator = "ECON, generado por scripts/docs/exportar_entregables.py"
    book.properties.lastModifiedBy = None
    book.properties.created = stamp
    book.properties.modified = stamp
    raw = io.BytesIO()
    # ExcelWriter directly: Workbook.save() would stamp `modified` with the wall clock.
    ExcelWriter(book, zipfile.ZipFile(raw, "w", zipfile.ZIP_DEFLATED, allowZip64=True)).save()
    source = zipfile.ZipFile(io.BytesIO(raw.getvalue()))
    fixed = io.BytesIO()
    with zipfile.ZipFile(fixed, "w", zipfile.ZIP_DEFLATED) as target:
        for name in source.namelist():
            info = zipfile.ZipInfo(name, date_time=ZIP_EPOCH)
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, source.read(name))
    return fixed.getvalue()


def unique(title: str, taken: set[str]) -> str:
    """Excel sheet titles are unique and at most 31 characters."""
    candidate = title[:31]
    suffix = 2
    while candidate in taken:
        candidate = f"{title[:28]} {suffix}"
        suffix += 1
    taken.add(candidate)
    return candidate


def build(documents: dict[Path, str], generated_on: date) -> dict[str, bytes]:
    """Every output file as bytes, keyed by path relative to the repository root."""
    detail = matrices.extract({path: documents[path] for path in MATRIX_SOURCES})
    by_slug = {table.spec.slug: table for table in detail}
    files: dict[str, bytes] = {}
    entries: list[dict] = []

    for deliverable in DELIVERABLES:
        spec = TableSpec(
            deliverable.stem.lower(),
            deliverable.sheet,
            deliverable.heading,
            requirement=deliverable.title,
            source=deliverable.document,
        )
        own = deliverable_table(spec, documents[deliverable.document], deliverable.document)
        sheets: list[tuple[Table, str]] = []
        taken: set[str] = set()
        sheets.append((own, unique(deliverable.sheet, taken)))
        selected = (
            detail
            if deliverable.extra_sheets == ("*",)
            else [by_slug[slug] for slug in deliverable.extra_sheets]
        )
        for table in selected:
            sheets.append((table, unique(table.spec.sheet, taken)))

        csv_path = OUTPUT_DIR / f"{deliverable.stem}.csv"
        data = csv_bytes(own)
        files[csv_path.as_posix()] = data
        entries.append(
            {
                "file": csv_path.as_posix(),
                "kind": "csv",
                "source": deliverable.document.as_posix(),
                "heading": own.section,
                "columns": own.header,
                "rows": len(own.rows),
                "sha256": sha256_bytes(data),
            }
        )
        for slug in deliverable.detail_slugs:
            table = by_slug[slug]
            detail_path = OUTPUT_DIR / deliverable.detail_csv
            detail_data = csv_bytes(table)
            files[detail_path.as_posix()] = detail_data
            entries.append(
                {
                    "file": detail_path.as_posix(),
                    "kind": "csv",
                    "source": table.spec.source.as_posix(),
                    "heading": table.section,
                    "columns": table.header,
                    "rows": len(table.rows),
                    "sha256": sha256_bytes(detail_data),
                }
            )
        book_path = OUTPUT_DIR / f"{deliverable.stem}.xlsx"
        book_data = workbook_bytes(sheets, deliverable.title, generated_on)
        files[book_path.as_posix()] = book_data
        entries.append(
            {
                "file": book_path.as_posix(),
                "kind": "xlsx",
                "source": deliverable.document.as_posix(),
                "heading": deliverable.title,
                "sheets": [title for _, title in sheets],
                "rows": sum(len(table.rows) for table, _ in sheets),
                "sha256": sha256_bytes(book_data),
            }
        )

    manifest = {
        "schema_version": 1,
        "generator": "scripts/docs/exportar_entregables.py",
        "reproduce": REPRODUCE,
        "check": f"{REPRODUCE} --check",
        "generated_on": generated_on.isoformat(),
        "csv_format": {
            "encoding": "utf-8-sig",
            "newline": "\\r\\n",
            "delimiter": ",",
            "note": (
                "BOM y CRLF para que Excel muestre los acentos al abrir el archivo con doble "
                "clic. El separador es la coma de RFC 4180: el CSV es para ingesta "
                "automatizada y el XLSX es el archivo para abrir en Excel."
            ),
        },
        "sources": [
            {"path": path.as_posix(), "sha256": sha256_bytes(text.encode("utf-8"))}
            for path, text in documents.items()
        ],
        "files": entries,
        "scope_notes": [
            "El Markdown de cada entregable es el origen editable; CSV y XLSX son copias.",
            "Cada archivo contiene la misma matriz que el entregable cuyo nombre lleva.",
            "Las hojas de detalle proceden de docs/ y coinciden con output/matrices.",
            "La RACI es una propuesta pendiente de validación empresarial (RF-03).",
            "Una correspondencia documentada no acredita identidad de registros entre plataformas.",
        ],
    }
    files[MANIFEST_JSON.as_posix()] = (
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n"
    ).encode("utf-8")
    files[MANIFEST_MD.as_posix()] = manifest_markdown(manifest).encode("utf-8")
    return files


def manifest_markdown(manifest: dict) -> str:
    lines = [
        "# Tablas exportadas de los entregables",
        "",
        f"Generado el {manifest['generated_on']} por `{manifest['generator']}` desde:",
        "",
    ]
    for source in manifest["sources"]:
        path = source["path"]
        link = Path(path).name if path.startswith(f"{OUTPUT_DIR.as_posix()}/") else f"../{path}"
        lines.append(f"- [`{path}`]({link}) (SHA-256 `{source['sha256']}`).")
    fmt = manifest["csv_format"]
    lines.extend(
        [
            "",
            "El Markdown de cada entregable sigue siendo el origen editable; estos archivos",
            "son copias derivadas para hojas de cálculo (RNF-02). No añaden filas ni valores.",
            "",
            f"**Formato CSV**: `{fmt['encoding']}` (UTF-8 con BOM), fin de línea `CRLF`, "
            f"separador `{fmt['delimiter']}`.",
            "",
            fmt["note"],
            "",
            "**Formato XLSX**: una hoja por tabla, fila de cabecera fija y filtrable, ancho de",
            "columna calculado, ajuste de texto y alineación superior.",
            "",
            "Reproducir y verificar:",
            "",
            "```sh",
            manifest["reproduce"],
            manifest["check"],
            "```",
            "",
            "| Archivo | Origen | Contenido | Filas | SHA-256 |",
            "| --- | --- | --- | --- | --- |",
        ]
    )
    for entry in manifest["files"]:
        name = Path(entry["file"]).name
        if entry["kind"] == "xlsx":
            content = f"{len(entry['sheets'])} hojas: {entry['sheets'][0]} …"
        else:
            content = f"{len(entry['columns'])} columnas: {entry['heading']}"
        lines.append(
            f"| [`{name}`]({name}) | `{Path(entry['source']).name}` | {content} | "
            f"{entry['rows']} | `{entry['sha256']}` |"
        )
    lines.extend(["", *(f"- {note}" for note in manifest["scope_notes"]), ""])
    return "\n".join(lines)


def stored_generated_on() -> date | None:
    path = ROOT / MANIFEST_JSON
    if not path.exists():
        return None
    try:
        return date.fromisoformat(json.loads(path.read_text(encoding="utf-8"))["generated_on"])
    except (KeyError, ValueError, json.JSONDecodeError):
        return None


def check(files: dict[str, bytes]) -> int:
    problems = []
    for relative, expected in files.items():
        path = ROOT / relative
        if not path.exists():
            problems.append(f"falta {relative}")
        elif path.read_bytes() != expected:
            problems.append(f"desactualizado {relative}")
    if problems:
        print("Tablas de ENTREGABLES no coinciden con el Markdown de origen:", file=sys.stderr)
        for problem in problems:
            print(f"  - {problem}", file=sys.stderr)
        print(f"Regenerar con: {REPRODUCE}", file=sys.stderr)
        return 1
    print(f"Entregables vigentes: {len(files) - 2} tablas y manifiestos coinciden.")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument(
        "--check", action="store_true", help="Compare the checked-in exports; write nothing"
    )
    args = parser.parse_args()
    sources = (*MATRIX_SOURCES, MAPPING_DOC, RACI_DOC)
    documents = {path: (ROOT / path).read_text(encoding="utf-8") for path in sources}
    today = datetime.now(UTC).date()
    generated_on = (stored_generated_on() or today) if args.check else today
    files = build(documents, generated_on)
    if args.check:
        return check(files)
    for relative, data in files.items():
        (ROOT / relative).write_bytes(data)
    print(
        f"Generado en {OUTPUT_DIR}: {len(files) - 2} tablas, manifest-tablas.json y "
        f"MANIFEST-TABLAS.md"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
